from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import fitz
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pptx import Presentation

app = FastAPI(title="Browser Agent File Parser", docs_url=None, redoc_url=None)
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
MAX_ZIP_UNPACKED = 100 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".csv", ".txt", ".md"}


@app.get("/health")
def health():
    return {"ok": True}


@app.post("/parse")
async def parse_file(file: UploadFile = File(...)):
    name = Path((file.filename or "upload").replace("\\", "/")).name
    extension = Path(name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(415, "Unsupported file type")
    data = await file.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(413, "File is larger than 50 MiB")
    try:
        text = extract_text(extension, data)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "File could not be parsed") from exc
    return {"name": name, "sizeBytes": len(data), "text": text[:MAX_TEXT_CHARS], "truncated": len(text) > MAX_TEXT_CHARS}


def check_archive(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_ZIP_UNPACKED:
            raise HTTPException(413, "Office document expands beyond the allowed size")
        if any(item.flag_bits & 0x1 for item in entries):
            raise HTTPException(422, "Encrypted office documents are not supported")


def extract_text(extension: str, data: bytes) -> str:
    if extension in {".txt", ".md"}:
        return data.decode("utf-8-sig", errors="replace")
    if extension == ".csv":
        rows = csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace")))
        return "\n".join("\t".join(cell[:2000] for cell in row[:100]) for _, row in zip(range(20_000), rows))
    if extension == ".pdf":
        document = fitz.open(stream=data, filetype="pdf")
        if document.page_count > 300:
            raise HTTPException(413, "PDF has more than 300 pages")
        return "\n\n".join(page.get_text("text") for page in document)
    if extension == ".docx":
        check_archive(data)
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables[:100]:
            parts.extend("\t".join(cell.text for cell in row.cells[:30]) for row in table.rows[:1000])
        return "\n".join(parts)
    if extension == ".xlsx":
        check_archive(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet in workbook.worksheets[:20]:
            parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(max_row=5000, max_col=50, values_only=True):
                parts.append("\t".join("" if value is None else str(value)[:1000] for value in row))
        workbook.close()
        return "\n".join(parts)
    if extension == ".pptx":
        check_archive(data)
        presentation = Presentation(io.BytesIO(data))
        if len(presentation.slides) > 200:
            raise HTTPException(413, "Presentation has more than 200 slides")
        parts = []
        for index, slide in enumerate(presentation.slides, start=1):
            parts.append(f"[Slide {index}]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text[:5000])
                if shape.has_table:
                    for row in shape.table.rows[:200]:
                        parts.append("\t".join(cell.text[:1000] for cell in row.cells[:30]))
        return "\n".join(parts)
    raise HTTPException(415, "Unsupported file type")
